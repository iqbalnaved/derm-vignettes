import re
import json
from docx import Document
from collections import Counter

# === Load DOCX ===
# file_path = "/mnt/d/Naved/Data/derm_vignette/Supplementary Text 3_10312025.docx"
# output_file = "/mnt/d/Naved/Data/derm_vignette/Supplementary_Text_3_10312025.json"

# file_path = "/mnt/d/Naved/Data/derm_vignette/Supplementary Text 3_11032025_v2.docx"
# output_file = "/mnt/d/Naved/Data/derm_vignette/Supplementary_Text_3_11032025_v2.json"

file_path = "/mnt/d/Naved/Data/derm_vignette/Supplementary Text 3_11042025_v1.docx"
output_file = "/mnt/d/Naved/Data/derm_vignette/Supplementary_Text_3_11042025_v1.json"

doc = Document(file_path)

vignettes = []

def clean_vignette(text):
    # remove the question line if present
    text = re.split(r"['‘’\"]?\s*What is the most likely diagnosis", text, flags=re.IGNORECASE)[0]
    text = re.sub(r"\s+", " ", text).strip()
    return text

# === Loop through tables ===
for t_i, table in enumerate(doc.tables):
    # Get the diagnosis header (e.g., “melanoma”, “nevus”)
    header_text = table.cell(0, 0).text.strip()
    
    # ---Supplementary_Text_3_11032025_v2.json and Supplementary_Text_3_10312025.json
    # diagnosis_match = re.search(r'["“]?\s*([A-Za-z\s]+?)\s*["”]?$', header_text)
    # diagnosis = diagnosis_match.group(1).strip() if diagnosis_match else "Unknown"

    # ---Supplementary_Text_3_11042025_v1.json-----
    # Extract everything after the plus sign
    plus_part = re.split(r'\+', header_text, maxsplit=1)[-1].strip()

    # Capture the text inside quotes (ignore abbreviation)
    diagnosis_match = re.search(r'["“”]\s*([A-Za-z\s]+?)\s*["“”]', plus_part)
    diagnosis = diagnosis_match.group(1).strip().title() if diagnosis_match else "Unknown"



    # Skip the first row (header)
    for r_i, row in enumerate(table.rows[1:], start=1):
        if len(row.cells) < 2:
            continue

        case_num = row.cells[0].text.strip()
        vignette_text = row.cells[1].text.strip()

        if not case_num or not vignette_text:
            continue

        vignette_text = clean_vignette(vignette_text)
        if vignette_text:
            vignettes.append({
                "case_number": int(case_num),
                "diagnosis": diagnosis,
                "vignette": vignette_text
            })

# === Save to JSON ===

with open(output_file, "w", encoding="utf-8") as f:
    json.dump(vignettes, f, indent=2, ensure_ascii=False)

print(f"✅ Extracted {len(vignettes)} vignettes and saved to '{output_file}'")

# === Frequency summary ===
counts = Counter(v["diagnosis"] for v in vignettes)
print("\n📊 Diagnosis frequencies:")
for diag, n in counts.items():
    print(f"  {diag}: {n}")

# Preview
for v in vignettes[:3]:
    print(f"\nCase #{v['case_number']} ({v['diagnosis']}):\n{v['vignette'][:200]}...\n")
